import asyncio
import os
import pandas as pd
from dotenv import load_dotenv
import aiohttp
import datetime as dt
import json
from urllib.parse import quote

from azure.identity.aio import ClientSecretCredential

# Load environment variables
load_dotenv()

TENANT_ID = os.getenv("TENANT_ID")
CLIENT_ID = os.getenv("CLIENT_ID")
CLIENT_SECRET = os.getenv("CLIENT_SECRET")
SITE_ID = os.getenv("SITE_ID")

# Custom API base URL (if different from Graph)
SHAREPOINT_API_BASE_URL = os.getenv("SHAREPOINT_API_BASE_URL", "https://graph.microsoft.com/v1.0")


async def get_bearer_token():
    """Get bearer token using ClientSecretCredential"""
    credential = ClientSecretCredential(
        tenant_id=TENANT_ID,
        client_id=CLIENT_ID,
        client_secret=CLIENT_SECRET
    )
    token_response = await credential.get_token("https://graph.microsoft.com/.default")
    return token_response.token


async def get_drive_id_from_site(base_api_url: str, site_id: str) -> str | None:
    """Get drive ID from site ID using REST API"""
    bearer_token = await get_bearer_token()
    if not bearer_token:
        print("❌ Failed to get bearer token")
        return None

    headers = {"Authorization": f"Bearer {bearer_token}"}

    async with aiohttp.ClientSession(headers=headers) as session:
        # Try to get drives from site
        short_site_id = site_id.split(',')[1] if ',' in site_id else site_id
        drives_url = f"{base_api_url}/sites/drives/{short_site_id}"

        try:
            async with session.get(drives_url) as response:
                if response.status == 200:
                    drives_content = await response.text()
                    drives_json = json.loads(drives_content)

                    # Try to parse drives array
                    if isinstance(drives_json, list):
                        drives_array = drives_json
                    elif "value" in drives_json:
                        drives_array = drives_json["value"]
                    elif "id" in drives_json:
                        return drives_json["id"]
                    else:
                        # Try lists fallback
                        pass

                    # Get first drive
                    if isinstance(drives_array, list):
                        for drive in drives_array:
                            if "id" in drive:
                                return drive["id"]
                else:
                    print(f"   Drives API failed: {response.status} - {await response.text()}")

        except Exception as e:
            print(f"   Error getting drives: {e}")

        # Fallback: Try getting lists to find Documents library
        try:
            lists_url = f"{base_api_url}/sites/{site_id}/lists"
            async with session.get(lists_url) as response:
                if response.status == 200:
                    lists_content = await response.text()
                    lists_json = json.loads(lists_content)

                    lists_array = lists_json.get("value", []) if isinstance(lists_json, dict) else lists_json

                    # Look for Documents library
                    for list_item in lists_array:
                        if list_item.get("displayName", "").lower() == "documents":
                            return site_id  # Use site_id as drive reference

                    return site_id  # Default fallback
                else:
                    return site_id  # Fallback
        except Exception as e:
            print(f"   Error getting lists: {e}")
            return site_id


async def find_file_in_path(base_api_url: str, drive_id: str, folder_path: str, file_name: str) -> str | None:
    """Find file in folder path using REST API"""
    bearer_token = await get_bearer_token()
    if not bearer_token:
        print("❌ Failed to get bearer token")
        return None

    headers = {"Authorization": f"Bearer {bearer_token}"}

    async with aiohttp.ClientSession(headers=headers) as session:
        # Use listContentFolder endpoint
        files_url = f"{base_api_url}/drives/listContentFolder/{drive_id}/{quote(folder_path)}"

        try:
            async with session.get(files_url) as response:
                if response.status == 200:
                    files_content = await response.text()
                    files_json = json.loads(files_content)

                    files = files_json if isinstance(files_json, list) else files_json.get("value", [])

                    for file in files:
                        if file.get("name", "").lower() == file_name.lower():
                            return file.get("id")
                else:
                    print(f"❌ Failed to list files: {response.status}")
                    error = await response.text()
                    print(f"   Error: {error[:300]}")
                    return None
        except Exception as e:
            print(f"   Error listing files: {e}")
            return None

    return None


async def download_excel_file(base_api_url: str, drive_id: str, file_id: str) -> bytes | None:
    """Download Excel file using REST API"""
    bearer_token = await get_bearer_token()
    if not bearer_token:
        print("❌ Failed to get bearer token")
        return None

    headers = {"Authorization": f"Bearer {bearer_token}"}

    async with aiohttp.ClientSession(headers=headers) as session:
        url = f"{base_api_url}/drives/downloadFile/{drive_id}/{file_id}"

        try:
            async with session.get(url) as response:
                if response.status == 200:
                    return await response.read()
                else:
                    print(f"❌ Failed to download file: {response.status}")
                    error = await response.text()
                    print(f"   Error: {error}")
                    return None
        except Exception as e:
            print(f"   Error downloading file: {e}")
            return None


def parse_excel_file(excel_bytes: bytes) -> pd.DataFrame | None:
    """Parse Excel file bytes to pandas DataFrame"""
    try:
        from io import BytesIO
        import openpyxl

        stream = BytesIO(excel_bytes)
        workbook = openpyxl.load_workbook(stream, data_only=True)

        worksheet = workbook.active
        if not worksheet:
            print("❌ No worksheets found in the Excel file")
            return None

        print(f"✅ Reading from worksheet: '{worksheet.title}'")

        # Get used range
        min_row = worksheet.min_row
        max_row = worksheet.max_row
        min_col = worksheet.min_column
        max_col = worksheet.max_column

        # Read all data
        data = []
        for row in range(min_row, max_row + 1):
            row_data = []
            for col in range(min_col, max_col + 1):
                cell_value = worksheet.cell(row=row, column=col).value
                row_data.append(str(cell_value) if cell_value is not None else "")
            data.append(row_data)

        if not data:
            print("  - The worksheet appears to be empty.")
            return None

        # Create DataFrame
        header = data[0]
        rows = data[1:]
        df = pd.DataFrame(rows, columns=header)
        df.replace("", None, inplace=True)
        print("✅ Successfully created pandas DataFrame.")
        return df

    except Exception as e:
        print(f"❌ Error parsing Excel file: {e}")
        return None


async def find_and_read_excel_file(site_id: str):
    print("\nAttempting to find and read 'CARTEIRA GRUPO ASSAÍ.xlsx'...")

    path_segments = [
        "Geral Alpargatas LLP",
        "19. Base RPA"
    ]
    file_name = "CARTEIRA GRUPO ASSAÍ.xlsx"
    folder_path = "/".join(path_segments)

    try:
        # Step 1: Get drive ID from site
        drive_id = await get_drive_id_from_site(SHAREPOINT_API_BASE_URL, site_id)
        if not drive_id:
            print("❌ Could not retrieve drive ID")
            return None, None, None

        print(f"✅ Retrieved Drive ID: {drive_id}")

        # Step 2: Find the file
        file_id = await find_file_in_path(SHAREPOINT_API_BASE_URL, drive_id, folder_path, file_name)
        if not file_id:
            print(f"❌ File '{file_name}' not found in folder '{folder_path}'")
            return None, None, None

        print(f"✅ Found file '{file_name}' with ID: {file_id}")

        # Step 3: Download the Excel file
        excel_bytes = await download_excel_file(SHAREPOINT_API_BASE_URL, drive_id, file_id)
        if not excel_bytes:
            print("❌ Failed to download Excel file")
            return None, None, None

        print(f"✅ Downloaded Excel file ({len(excel_bytes)} bytes)")

        # Step 4: Parse the Excel file
        df = parse_excel_file(excel_bytes)
        if df is not None:
            return df, drive_id, file_id
        else:
            return None, None, None

    except Exception as e:
        print(f"❌ Error retrieving Excel data: {e}")
        return None, None, None



# ---------------- UPDATE FUNCTION ----------------
async def update_protocol_rows(drive_id: str, file_id: str, protocol_data_list: list[dict]):
    """
    Updates rows in Excel with protocol numbers where chave_pedido_loja + CARRO matches.
    
    Args:
        drive_id: SharePoint drive ID
        file_id: Excel file ID
        protocol_data_list: List of dicts with {chave, carro, protocol}
    
    Updates column 'PROTOCOLO DA SOLICITAÇÃO' (column BL) where:
        - chave_pedido_loja = 'Nº Pedido Cliente' + '-' + first_part_of('CÓD LOJA')
        - CARRO matches
    """
    print("\n🔄 Starting protocol update process...")
    print(f"  - Protocol data to update: {protocol_data_list}")

    try:
        # Get bearer token
        bearer_token = await get_bearer_token()
        if not bearer_token:
            print("❌ Failed to get bearer token")
            return

        headers = {"Authorization": f"Bearer {bearer_token}"}

        async with aiohttp.ClientSession(headers=headers) as session:
            # Get worksheets
            worksheets_url = f"{SHAREPOINT_API_BASE_URL}/drives/{drive_id}/items/{file_id}/workbook/worksheets"
            async with session.get(worksheets_url) as response:
                if response.status != 200:
                    print(f"❌ Failed to get worksheets: {response.status}")
                    return
                worksheets_data = await response.json()

            worksheets = worksheets_data.get("value", [])
            if not worksheets:
                print("❌ No worksheets found.")
                return

            first_worksheet = worksheets[0]
            worksheet_id = first_worksheet["id"]
            print(f"  - Target worksheet: {first_worksheet['name']}")

            # Get used range
            used_range_url = f"{SHAREPOINT_API_BASE_URL}/drives/{drive_id}/items/{file_id}/workbook/worksheets/{worksheet_id}/usedRange"
            async with session.get(used_range_url) as response:
                if response.status != 200:
                    print(f"❌ Failed to get used range: {response.status}")
                    return
                used_range_data = await response.json()

            if "values" not in used_range_data:
                print("❌ No data found to update.")
                return

            if "address" not in used_range_data:
                print("❌ Could not determine range address for update.")
                return

            values = used_range_data["values"]
            header, data = values[0], values[1:]

            # Find column indices
            col_indices = {}
            for idx, col_name in enumerate(header):
                col_name_lower = str(col_name).lower().strip() if col_name else ""
                if 'nº pedido cliente' in col_name_lower or 'pedido cliente' in col_name_lower:
                    col_indices['pedido_cliente'] = idx
                elif 'cód loja' in col_name_lower or 'cod loja' in col_name_lower:
                    col_indices['cod_loja'] = idx
                elif 'carro' in col_name_lower:
                    col_indices['carro'] = idx
                elif 'protocolo' in col_name_lower and 'solicitação' in col_name_lower:
                    col_indices['protocolo'] = idx
            
            # Verify required columns exist
            required_cols = ['pedido_cliente', 'cod_loja', 'carro', 'protocolo']
            missing_cols = [col for col in required_cols if col not in col_indices]
            if missing_cols:
                print(f"❌ Missing required columns: {missing_cols}")
                return
            
            # Update matching rows
            updated_count = 0
            for row in data:
                if len(row) <= max(col_indices.values()):
                    # Extend row if needed
                    row.extend([""] * (max(col_indices.values()) + 1 - len(row)))
                
                # Build chave_pedido_loja from row data
                pedido_val = str(row[col_indices['pedido_cliente']]).strip() if row[col_indices['pedido_cliente']] else ""
                loja_val = str(row[col_indices['cod_loja']]).strip() if row[col_indices['cod_loja']] else ""
                carro_val = str(row[col_indices['carro']]).strip() if row[col_indices['carro']] else ""
                
                # Extract first part of loja (before '-')
                loja_first_part = loja_val.split('-')[0] if loja_val else ""
                row_chave = f"{pedido_val}-{loja_first_part}" if pedido_val and loja_first_part else ""
                
                # Check if this row matches any protocol data
                for protocol_item in protocol_data_list:
                    item_chave = str(protocol_item.get('chave', '')).strip()
                    item_carro = str(protocol_item.get('carro', '')).strip()
                    item_protocol = str(protocol_item.get('protocol', '')).strip()
                    
                    if row_chave == item_chave and carro_val == item_carro:
                        # Update protocol column
                        row[col_indices['protocolo']] = item_protocol
                        updated_count += 1
                        break

            updated_values = [header] + data

            if updated_count == 0:
                print("  - No matching rows found to update.")
                return

            # Direct REST call to update
            target_address = used_range_data["address"]
            print(f"  - Updating range: {target_address}")

            endpoint = f"{SHAREPOINT_API_BASE_URL}/drives/{drive_id}/items/{file_id}/workbook/worksheets/{worksheet_id}/range(address='{target_address}')"

            async with session.patch(
                endpoint,
                headers={"Content-Type": "application/json"},
                json={"values": updated_values},
            ) as resp:
                if resp.status == 200:
                    print(f"✅ Successfully updated {updated_count} protocol rows in Excel.")
                else:
                    text = await resp.text()
                    print(f"❌ Protocol update failed ({resp.status}): {text}")

    except Exception as ex:
        print(f"❌ Unexpected error during protocol update: {ex}")
        import traceback
        traceback.print_exc()


# Keep old function name for backward compatibility (redirects to new function)
async def update_excel_rows(drive_id: str, file_id: str, lookup_values: list[str]):
    """Legacy function - redirects to update_protocol_rows with old behavior"""
    print("⚠️ Using legacy update_excel_rows function")
    # This maintains old behavior if called elsewhere
    protocol_data = [{"chave": val, "carro": "", "protocol": "Não Encontrado"} for val in lookup_values]
    await update_protocol_rows(drive_id, file_id, protocol_data)


# ---------------- UPDATE AGENDA COLUMNS (NEW FOR PEGAR RETORNO) ----------------
async def update_agenda_columns(drive_id: str, file_id: str, agenda_data_list: list[dict]):
    """
    Updates the 'AGENDA CONFIRMADA', 'PROTOCOLO AGENDA', and 'HORÁRIO' columns in SharePoint Excel with extracted data.
    
    Args:
        drive_id: SharePoint drive ID
        file_id: Excel file ID
        agenda_data_list: List of dicts with structure:
            [{"chave": "12345-67", "protocol": "P001", "agenda_confirmada": "12/08/2025", "protocolo_agenda": "Agendamento12345", "horario": "14:30"}, ...]
    
    Matching logic:
        - Matches rows where chave_pedido_loja == chave AND protocol == protocol
        - Updates the respective columns with the provided values
    """
    try:
        print("\n🔄 Starting agenda columns update process...")
        print(f"  - Processing {len(agenda_data_list)} agenda updates")

        # Get bearer token
        bearer_token = await get_bearer_token()
        if not bearer_token:
            print("❌ Failed to get bearer token")
            return

        headers = {"Authorization": f"Bearer {bearer_token}"}

        async with aiohttp.ClientSession(headers=headers) as session:
            # Get worksheets
            worksheets_url = f"{SHAREPOINT_API_BASE_URL}/drives/{drive_id}/items/{file_id}/workbook/worksheets"
            async with session.get(worksheets_url) as response:
                if response.status != 200:
                    print(f"❌ Failed to get worksheets: {response.status}")
                    return
                worksheets_data = await response.json()

            worksheets = worksheets_data.get("value", [])
            if not worksheets:
                print("❌ No worksheets found.")
                return

            first_worksheet = worksheets[0]
            worksheet_id = first_worksheet["id"]
            print(f"  - Target worksheet: {first_worksheet['name']}")

            # Get used range
            used_range_url = f"{SHAREPOINT_API_BASE_URL}/drives/{drive_id}/items/{file_id}/workbook/worksheets/{worksheet_id}/usedRange"
            async with session.get(used_range_url) as response:
                if response.status != 200:
                    print(f"❌ Failed to get used range: {response.status}")
                    return
                used_range_data = await response.json()

            if "values" not in used_range_data:
                print("❌ No data found to update.")
                return

            if "address" not in used_range_data:
                print("❌ Could not determine range address for update.")
                return

            values = used_range_data["values"]
            header, data = values[0], values[1:]

            # Convert any datetime objects to strings to avoid JSON serialization errors
            for row in data:
                for idx, val in enumerate(row):
                    if isinstance(val, dt.datetime):
                        row[idx] = str(val)

            # Find column indices
            col_indices = {}
            for idx, col_name in enumerate(header):
                col_name_lower = str(col_name).lower().strip() if col_name else ""
                if 'nº pedido cliente' in col_name_lower or 'pedido cliente' in col_name_lower:
                    col_indices['pedido_cliente'] = idx
                elif 'cód loja' in col_name_lower or 'cod loja' in col_name_lower:
                    col_indices['cod_loja'] = idx
                elif 'protocolo' in col_name_lower and 'solicitação' in col_name_lower:
                    col_indices['protocolo'] = idx
                elif 'agenda confirmada' in col_name_lower:
                    col_indices['agenda_confirmada'] = idx
                elif 'protocolo agenda' in col_name_lower:
                    col_indices['protocolo_agenda'] = idx
                elif 'horário' in col_name_lower or 'horario' in col_name_lower:
                    col_indices['horario'] = idx
            
            # Verify required columns exist
            required_cols = ['pedido_cliente', 'cod_loja', 'protocolo', 'agenda_confirmada', 'protocolo_agenda', 'horario']
            missing_cols = [col for col in required_cols if col not in col_indices]
            if missing_cols:
                print(f"❌ Missing required columns: {missing_cols}")
                return
            
            # Update matching rows
            updated_count = 0
            for row in data:
                if len(row) <= max(col_indices.values()):
                    # Extend row if needed
                    row.extend([""] * (max(col_indices.values()) + 1 - len(row)))
                
                # Build chave_pedido_loja from row data
                pedido_val = str(row[col_indices['pedido_cliente']]).strip() if row[col_indices['pedido_cliente']] else ""
                loja_val = str(row[col_indices['cod_loja']]).strip() if row[col_indices['cod_loja']] else ""
                protocol_val = str(row[col_indices['protocolo']]).strip() if row[col_indices['protocolo']] else ""
                
                # Extract first part of loja (before '-')
                loja_first_part = loja_val.split('-')[0] if loja_val else ""
                row_chave = f"{pedido_val}-{loja_first_part}" if pedido_val and loja_first_part else ""
                
                # Check if this row matches any agenda data
                for agenda_item in agenda_data_list:
                    item_chave = str(agenda_item.get('chave', '')).strip()
                    item_protocol = str(agenda_item.get('protocol', '')).strip()
                    item_agenda_confirmada = str(agenda_item.get('agenda_confirmada', '')).strip()
                    item_protocolo_agenda = str(agenda_item.get('protocolo_agenda', '')).strip()
                    item_horario = str(agenda_item.get('horario', '')).strip()
                    
                    if row_chave == item_chave and protocol_val == item_protocol:
                        # Update the three columns
                        row[col_indices['agenda_confirmada']] = item_agenda_confirmada
                        row[col_indices['protocolo_agenda']] = f"'{item_protocolo_agenda}"  # Prefix with ' to ensure text format in Excel
                        row[col_indices['horario']] = item_horario
                        updated_count += 1
                        break

            updated_values = [header] + data

            if updated_count == 0:
                print("  - No matching rows found to update.")
                return

            # Direct REST call to update
            target_address = used_range_data["address"]
            print(f"  - Updating range: {target_address}")

            endpoint = f"{SHAREPOINT_API_BASE_URL}/drives/{drive_id}/items/{file_id}/workbook/worksheets/{worksheet_id}/range(address='{target_address}')"

            async with session.patch(
                endpoint,
                headers={"Content-Type": "application/json"},
                json={"values": updated_values},
            ) as resp:
                if resp.status == 200:
                    print(f"✅ Successfully updated {updated_count} agenda columns in Excel.")
                else:
                    text = await resp.text()
                    print(f"❌ Agenda update failed ({resp.status}): {text}")

    except Exception as ex:
        print(f"❌ Unexpected error during agenda update: {ex}")
        import traceback
        traceback.print_exc()



# ---------------- ASYNC WRAPPER FOR PROTOCOL UPDATE ----------------
async def update_protocol_async(drive_id: str, file_id: str, protocol_data_list: list[dict]):
    """
    Async wrapper to update protocols in SharePoint Excel.
    Can be called from a background thread.
    
    Args:
        drive_id: SharePoint drive ID
        file_id: Excel file ID  
        protocol_data_list: List of dicts with {chave, carro, protocol}
    """
    try:
        await update_protocol_rows(drive_id, file_id, protocol_data_list)
        # print("✅ Protocol update completed successfully")
    except Exception as e:
        print(f"❌ Protocol update failed: {e}")
        import traceback
        traceback.print_exc()


# ---------------- ASYNC WRAPPER FOR AGENDA UPDATE ----------------
async def update_agenda_async(drive_id: str, file_id: str, agenda_data_list: list[dict]):
    """
    Async wrapper to update agenda columns in SharePoint Excel.
    Can be called from a background thread.
    
    Args:
        drive_id: SharePoint drive ID
        file_id: Excel file ID
        agenda_data_list: List of dicts with {chave, protocol, agenda_confirmada, protocolo_agenda, horario}
    """
    try:
        await update_agenda_columns(drive_id, file_id, agenda_data_list)
        print("✅ Agenda update completed successfully")
    except Exception as e:
        print(f"❌ Agenda update failed: {e}")
        import traceback
        traceback.print_exc()


# ---------------- MAIN ----------------
async def main():
    site_id = SITE_ID

    try:
        print(f"Connecting to SharePoint site: {site_id}")

        df, drive_id, file_id = await find_and_read_excel_file(site_id)

        if df is not None and not df.empty:
            # print(f"\n--- DataFrame Info ---")
            # print(f"Shape: {df.shape}")
            # print(df.head(2))

            # 🔹 Example usage: Update protocols
            # protocol_data = [{"chave": "12345-67", "carro": "CARRO1", "protocol": "98765"}]
            # await update_protocol_rows(drive_id, file_id, protocol_data)

            return df, drive_id, file_id 

        else:
            print("⚠️ Could not read any data.")

    except Exception as ex:
        print(f"❌ Unexpected error: {ex}")


# if __name__ == "__main__":
#     asyncio.run(main())